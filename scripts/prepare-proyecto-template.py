"""Build the editable UPS project template; no personal document is bundled.
Run with Python. The institutional crest comes from the supplied UPS guide.
"""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import base64
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(12); s.font.color.rgb=RGBColor(0,0,0)
 pf=s.paragraph_format; pf.space_before=Pt(0); pf.space_after=Pt(0); pf.line_spacing=2
 pf.first_line_indent=Inches(.5) if name=='Normal' else Inches(0)
 for element in list(s.element.iter()):
  if element.tag==qn('w:rFonts'):
   for attr in list(element.attrib):
    if 'theme' in attr.lower():del element.attrib[attr]
  if element.tag in [qn('w:pBdr'),qn('w:spacing')] and element.getparent().tag==qn('w:rPr'):element.getparent().remove(element)
 for border in s.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
 if name!='Normal': s.font.bold=True; pf.keep_with_next=True
 if name in ['Heading 1','Title']: pf.alignment=WD_ALIGN_PARAGRAPH.CENTER
 else: pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
for name in ['Cover','Table Text','Guide','Reference']:
 s=doc.styles.add_style(name,1); s.base_style=doc.styles['Normal']; s.paragraph_format.first_line_indent=Inches(0)
doc.styles['Cover'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Cover'].paragraph_format.line_spacing=1
doc.styles['Cover'].paragraph_format.space_after=Pt(5)
doc.styles['Table Text'].paragraph_format.line_spacing=1
doc.styles['Table Text'].paragraph_format.space_after=Pt(4)
doc.styles['Reference'].paragraph_format.left_indent=Inches(.5)
doc.styles['Reference'].paragraph_format.first_line_indent=Inches(-.5)

def setup(s,landscape=False):
 s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 s.page_width=Inches(11 if landscape else 8.5); s.page_height=Inches(8.5 if landscape else 11)
 s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
 s.header_distance=s.footer_distance=Inches(.5)
setup(doc.sections[0])
p=doc.sections[0].header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Inches(0)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
f=OxmlElement('w:updateFields');f.set(qn('w:val'),'false');doc.settings.element.append(f)

def p(text='',style=None):return doc.add_paragraph(text,style)
def cover(text,bold=False,size=Pt(12),space_before=Pt(0),space_after=Pt(4),line_spacing=1):
 par=p(text,'Cover')
 par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 par.paragraph_format.space_before=space_before
 par.paragraph_format.space_after=space_after
 par.paragraph_format.line_spacing=line_spacing
 for r in par.runs:
  r.bold=bold
  r.font.name='Times New Roman'
  r.font.size=size
  r.font.color.rgb=RGBColor(0,0,0)
 return par

index_titles=['Introducción','Identificación del problema','Delimitación del problema','Objetivos del proyecto','Metas del proyecto','Recursos','Cronograma pormenorizado','Cronograma por horas','Referencias','Anexos','Anexo A','Anexo B','Anexo C','Anexo D']

def heading(text,level=1,new=False):
 par=doc.add_heading(text,level);par.paragraph_format.page_break_before=new
 if text in index_titles:
  idx=index_titles.index(text)+1
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(idx));start.set(qn('w:name'),'section_'+str(idx))
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(idx))
  par._p.insert(0,start);par._p.append(end)
 return par

def paragraphs(key,style=None):
 p('{#'+key+'}',style);p('{text}',style);p('{/'+key+'}',style)

def table(headers,keys,widths):
 t=doc.add_table(rows=2, cols=len(headers)); t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,(h,k) in enumerate(zip(headers,keys)):
  for row,text in [(t.rows[0],h),(t.rows[1],k)]:
   c=row.cells[i];c.width=Inches(widths[i]);c.text=text
   for par in c.paragraphs:par.style='Table Text'
  for r in t.rows[0].cells[i].paragraphs[0].runs:r.bold=True
 trpr=t.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
 t._tbl.tblPr.append(borders)
 return t

cover('UNIVERSIDAD DE EL SALVADOR',bold=True,size=Pt(13.5),space_after=Pt(3))
cover('FACULTAD MULTIDISCIPLINARIA ORIENTAL',bold=True,size=Pt(12),space_after=Pt(3))
cover('DEPARTAMENTO: {departamento}',bold=False,size=Pt(11),space_after=Pt(2))
p('{#tieneSeccion}','Cover')
cover('SECCIÓN: {seccion}',bold=False,size=Pt(11),space_after=Pt(2))
p('{/tieneSeccion}','Cover')
cover('CARRERA: {carrera}',bold=False,size=Pt(11),space_after=Pt(8))
with ZipFile(ROOT/'documentos/procesos/proyecto.docx') as z:logo=z.read('word/media/image1.png')
par=p('','Cover')
par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.paragraph_format.space_before=Pt(4)
par.paragraph_format.space_after=Pt(12)
par.add_run().add_picture(BytesIO(logo),height=Inches(1.2))
cover('PROYECTO DE SERVICIO SOCIAL',bold=True,size=Pt(12.5),space_after=Pt(6))
par=p('{titulo}','Title')
par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.paragraph_format.line_spacing=1.15
par.paragraph_format.space_before=Pt(2)
par.paragraph_format.space_after=Pt(14)
for r in par.runs:
 r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(13.5);r.font.color.rgb=RGBColor(0,0,0)
cover('NÚMERO DE PROYECTO: {numero}',bold=False,size=Pt(10.5),space_after=Pt(2))
cover('Lugar de ejecución: {institucion} — {direccionInstitucion}',bold=False,size=Pt(10.5),space_after=Pt(2))
cover('Período de ejecución: Del {inicio} al {fin}',bold=False,size=Pt(10.5),space_after=Pt(10))
cover('PRESENTADO POR:',bold=True,size=Pt(11),space_after=Pt(3))
p('{#autores}','Cover')
cover('{nombre}    (Carné: {carnet})',bold=False,size=Pt(10.5),space_after=Pt(2))
p('{/autores}','Cover')
cover('DOCENTE TUTOR:',bold=True,size=Pt(11),space_before=Pt(6),space_after=Pt(2))
cover('{tutor}',bold=False,size=Pt(11),space_after=Pt(14))
cover('________________________________________\nFirma y sello del docente tutor',bold=False,size=Pt(9.5),space_after=Pt(12))
cover('UNIDAD DE PROYECCIÓN SOCIAL',bold=True,size=Pt(10.5),space_after=Pt(2))
cover('Ciudad Universitaria de Oriente, El Salvador',bold=False,size=Pt(10),space_after=Pt(2))
cover('{entrega}',bold=False,size=Pt(10),space_after=Pt(0))

heading('Índice',new=True)
for idx,title in enumerate(index_titles,1):
 if title=='Referencias':p('{#tieneReferencias}')
 par=p('', 'Guide')
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'section_'+str(idx))
 run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);par._p.append(link)
 if title=='Referencias':p('{/tieneReferencias}')
for title,key in [('Introducción','introduccion'),('Identificación del problema','problema'),('Delimitación del problema','delimitacion')]:
 heading(title,new=True);paragraphs(key)
heading('Objetivos del proyecto',new=True)
heading('Objetivo general',2);paragraphs('objetivo')
heading('Objetivos específicos',2);paragraphs('especificos')
heading('Metas del proyecto',new=True);paragraphs('metas')
heading('Recursos',new=True)
for title,key in [('Recursos humanos','humanos'),('Recursos materiales','materiales'),('Recursos financieros','financieros')]:heading(title,2);paragraphs(key)

setup(doc.add_section(WD_SECTION_START.NEW_PAGE),True)
heading('Cronograma pormenorizado')
p('Tabla 1','Guide').runs[0].bold=True
p('Actividades por período','Guide').runs[0].italic=True
p('Tutor: {tutor}\nEstudiantes: {participantesNombres}','Guide')
p('{@cronogramaDetalle}')
heading('Cronograma por horas',new=True)
p('Tabla 2','Guide').runs[0].bold=True;p('Distribución de horas por período','Guide').runs[0].italic=True
p('{@cronogramaHoras}')
p('Horas requeridas por cada alumno: {totalHoras}','Guide')
p('Nota. Cada estudiante tiene su propio cronograma. Las actividades se programan en el orden indicado, según su disponibilidad diaria y dentro del período del proyecto.','Guide')
setup(doc.add_section(WD_SECTION_START.NEW_PAGE))
p('{#tieneReferencias}')
heading('Referencias');paragraphs('referencias','Reference')
p('{/tieneReferencias}')
p('{#tieneReferencias}')
par=p('');par.paragraph_format.page_break_before=True
p('{/tieneReferencias}')
heading('Anexos',new=False)
p('Organiza los anexos en este orden: A. Documentos de inscripción; B. Respaldos del proyecto; C. Cuadro resumen; D. Solicitud de la institución. Sustituye las hojas de indicaciones por los documentos correspondientes antes de entregar.','Guide')
heading('Anexo A',new=True);heading('Documentos de inscripción',2)
p('Adjunta el expediente de todos los estudiantes, firmado y sellado, en este orden:','Guide')
for item in ['1. Carta de solicitud de la institución.','2. Formulario n.º 1 de cada estudiante.','3. Carta n.º 1 de asignación del tutor.','4. Carta n.º 2 de aprobación del plan de cada estudiante.','5. Constancia del 60 % de materias aprobadas de cada estudiante.']:p(item,'Guide')
p('Retira esta hoja cuando hayas incorporado los documentos.','Guide')
heading('Anexo B',new=True);heading('Respaldos del proyecto',2)
p('Adjunta los mapas, fotografías u otros respaldos que correspondan. Identifica cada uno con título y fuente. Si no corresponde incluir respaldos, indícalo con tu tutor.','Guide')
paragraphs('anexos','Guide')
p('Retira estas indicaciones al incorporar los respaldos. Las fuentes citadas en el texto se registran en Referencias.','Guide')
heading('Anexo C',new=True);heading('Cuadro resumen del proyecto',2)
# Label-value paragraphs stay readable even for very long author lists or objectives.
for label,key in [('Proyecto','titulo'),('Número de proyecto','numero'),('Estudiantes','participantes'),('Carrera','carrera'),('Institución y ubicación','institucionDireccion'),('Área de conocimiento','area'),('Población beneficiada','beneficiarios'),('Período de ejecución','periodo'),('Docente tutor','tutor')]:
 par=p('', 'Guide');par.add_run(label+': ').bold=True;par.add_run('{'+key+'}')
for title,key in [('Objetivo general','objetivo'),('Objetivos específicos','especificos'),('Recursos humanos','humanos'),('Recursos materiales','materiales'),('Presupuesto y procedencia de fondos','financieros')]:heading(title,2);paragraphs(key,'Guide')
heading('Anexo D',new=True);heading('Solicitud de la institución',2)
p('Adjunta la solicitud emitida por la institución que requiere a los estudiantes, con firma y sello de su responsable.','Guide')
p('La guía de proyecto también la enumera al final de los anexos. Si ya está en el expediente del Anexo A, confirma con coordinación si requiere otra copia o una referencia a esa página.','Guide')
p('Sustituye esta hoja por el documento o por la referencia que coordinación indique.','Guide')
# No source personal metadata, signatures, seals, example prose or identifiers.
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='Proyecto de servicio social';doc.core_properties.comments=''
out=ROOT/'documentos/plantillas/proyecto.docx';doc.save(out)
encoded=base64.b64encode(out.read_bytes()).decode()
(ROOT/'js/proyecto-plantilla.js').write_text('// Generated by scripts/prepare-proyecto-template.py\nwindow.SIGPSProyectoPlantilla="'+encoded+'";\n')
print(out)
