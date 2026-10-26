"""Development-only template authoring. The published portal uses JS and embedded DOCX files."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import base64
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
source=R/'Plantillas/3. DOCUMENTO PARA PRESENTAR LA MEMORIA DE SERVICIO SOCIAL.docx'
d=Document(source); originals=list(d.element.body)
# Isolate institutional styles so narrative formatting cannot alter original forms.
source_styles=list(d.styles.element.findall(qn('w:style')))
style_ids={el.get(qn('w:styleId')):'AnexoOriginal_'+el.get(qn('w:styleId')) for el in source_styles}
for el in source_styles:
 clone=deepcopy(el);clone.set(qn('w:styleId'),style_ids[el.get(qn('w:styleId'))]);clone.attrib.pop(qn('w:default'),None)
 for node in clone.iter():
  if node.tag in [qn('w:basedOn'),qn('w:next'),qn('w:link')]:
   val=node.get(qn('w:val'))
   if val in style_ids:node.set(qn('w:val'),style_ids[val])
 name=clone.find(qn('w:name'))
 if name is not None:name.set(qn('w:val'),'Anexo original '+name.get(qn('w:val')))
 d.styles.element.append(clone)
for el in list(d.element.body):
 if el.tag!=qn('w:sectPr'):d.element.body.remove(el)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Cambria';s.font.size=Pt(12);s.font.color.rgb=RGBColor(0,0,0)
 pf=s.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.5;pf.first_line_indent=Pt(0)
 pf.keep_with_next=name!='Normal';s.font.bold=name!='Normal'
 for el in list(s.element.iter()):
  if el.tag==qn('w:rFonts'):
   for a in list(el.attrib):
    if 'theme' in a.lower():del el.attrib[a]
  if el.tag==qn('w:pBdr'):el.getparent().remove(el)
 d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
s=d.sections[0];s.top_margin=s.bottom_margin=s.right_margin=Cm(2);s.left_margin=Cm(3)
# Original institutional header including its anchored crest, copied verbatim.
for i in [73,74]:d.element.body.insert(len(d.element.body)-1,deepcopy(originals[i]))
def p(text='',size=12,bold=False,center=False,space=6,line=1.15):
 q=d.add_paragraph(text);q.paragraph_format.space_after=Pt(space);q.paragraph_format.line_spacing=line;q.paragraph_format.first_line_indent=Pt(0)
 q.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
 for r in q.runs:r.font.name='Cambria';r.font.size=Pt(size);r.bold=bold
 return q
def head(text,new=True,mark=None):
 q=d.add_heading(text,1);q.paragraph_format.page_break_before=new
 if mark:
  a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(mark));a.set(qn('w:name'),'memoria_'+str(mark));b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(mark));q._p.insert(0,a);q._p.append(b)
 return q
def paras(key):
 d.add_paragraph('{#'+key+'}');d.add_paragraph('{text}');d.add_paragraph('{/'+key+'}')
p('DEPARTAMENTO: {departamento}',11,center=True,space=2)
p('SECCIÓN: {seccion}',11,center=True,space=20)
q=p('MEMORIA DE SERVICIO SOCIAL',14,True,True,12);q.style=d.styles['Title'];q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('{nombreMemoria}',13,True,True,16)
for text in ['No. de proyecto: {numero}','Ubicación: {institucion} — {direccionInstitucion}','Fecha de inicio: {inicio}','Fecha de finalización: {fin}']:p(text,11,space=8)
p('PRESENTADO POR',12,True,space=4)
t=d.add_table(rows=2,cols=2);t.autofit=False;t.columns[0].width=Cm(11.5);t.columns[1].width=Cm(5)
for row,texts in zip(t.rows,[['Nombre completo','Carnet'],['{#autores}{nombre}','{carnet}{/autores}']]):
 for c,text in zip(row.cells,texts):
  c.text=text
  for q in c.paragraphs:
   q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1
   for r in q.runs:r.font.size=Pt(11)
p('Carrera de todos los integrantes: {carrera}',11,space=14)
p('Docente tutor: {tutor}',11,space=24)
p('____________________________________________\nFirma y sello del docente tutor',11,center=True,space=14)
p('UNIDAD RESPONSABLE',11,True,space=2)
p('Unidad de Proyección Social y Subunidades de Proyección Social',11,space=12)
p('Fecha de entrega\nCiudad Universitaria de Oriente, {entrega}',11,space=0)
sections=[('Introducción','introduccion'),('Desarrollo','desarrollo'),('Resultados','resultados'),('Conclusiones','conclusiones'),('Recomendaciones','recomendaciones')]
head('Índice')
for i,(label,key) in enumerate(sections+[('Cuadro resumen','resumen'),('Anexos','anexos')],1):
 q=p('');link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'memoria_'+str(i));r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=label;r.append(tx);link.append(r);q._p.append(link)
for i,(label,key) in enumerate(sections,1):head(label,mark=i);paras(key)
head('Cuadro resumen',mark=6)
p('UNIVERSIDAD DE EL SALVADOR\nFACULTAD MULTIDISCIPLINARIA ORIENTAL\nUNIDAD DE PROYECCIÓN SOCIAL',11,True,center=True)
p('Memoria de servicio social · Ministerio de Educación',11,center=True)
for label,key in [('Nombre','nombreMemoria'),('Área de conocimiento','area'),('Cantidad de estudiantes','cantidad'),('Carrera','carrera'),('Fecha de inicio','inicio'),('Fecha de finalización','fin')]:p(label+': {'+key+'}',11,space=3)
# Retain the original six-column table geometry and borders, with automatic row height.
tbl=deepcopy(originals[151]);d.element.body.insert(len(d.element.body)-1,tbl);t=Table(tbl,d._body)
headers=['Población beneficiada (nombre y número)','Lugar donde se desarrolló','Objetivos','Logros','Recursos humanos, materiales y financieros','Procedencia de fondos']
keys=['{beneficiarios}','{institucion}\n{direccionInstitucion}','{objetivo}','{resultadosResumen}','{recursosResumen}','{procedenciaFondos}']
for ri,(row,texts) in enumerate(zip(t.rows,[headers,keys])):
 for c,text in zip(row.cells,texts):
  c.text=text
  for q in c.paragraphs:
   q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1;q.paragraph_format.keep_with_next=False
   for r in q.runs:r.font.name='Cambria';r.font.size=Pt(10);r.bold=ri==0
 for height in row._tr.findall('.//'+qn('w:trHeight')):height.getparent().remove(height)
 for cant in row._tr.findall('.//'+qn('w:cantSplit')):cant.getparent().remove(cant)
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
# Each instruction page is followed by the corresponding original blank form.
# These are preparation sheets: the completed, signed originals replace them for delivery.
form_ranges=[(177,197),(198,216),(216,255),(255,279),(288,312),(316,337),(341,368),(368,384),(386,397),(401,420)]
expected=['Formulario N.º 3','Formulario N.º 4','Formulario N.º 5','Carta N.º 3','Carta N.º 4','Carta N.º 5','Carta N° 6','Formulario N.º 6','Formulario Nº 7','Formulario Nº 8']
def append_original(start,end,label):
 assert ''.join(x.text or '' for x in originals[start].iter(qn('w:t'))).startswith(label), 'Source layout changed: '+label
 for offset,original in enumerate(originals[start:end]):
  el=deepcopy(original)
  for node in list(el.iter()):
   if node.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd'),qn('w:lastRenderedPageBreak')]:node.getparent().remove(node)
   elif node.tag in [qn('w:pStyle'),qn('w:rStyle'),qn('w:tblStyle')]:
    val=node.get(qn('w:val'))
    if val in style_ids:node.set(qn('w:val'),style_ids[val])
  for paragraph in el.iter(qn('w:p')):
   pr=paragraph.find(qn('w:pPr'))
   if pr is None:pr=OxmlElement('w:pPr');paragraph.insert(0,pr)
   if pr.find(qn('w:pStyle')) is None:
    st=OxmlElement('w:pStyle');st.set(qn('w:val'),style_ids['Normal']);pr.insert(0,st)
  if offset==0:
   pr=el.find(qn('w:pPr'));br=OxmlElement('w:pageBreakBefore');pr.append(br)
  d.element.body.insert(len(d.element.body)-1,el)
for i in range(1,14):
 q=head('HOJA DE INDICACIONES',mark=7 if i==1 else None)
 q.paragraph_format.space_before=Pt(60);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p('Anexo '+str(i),11,center=True,space=18)
 q=p('{anexo'+str(i)+'Nombre}',18,True,True,24);q.paragraph_format.keep_with_next=True
 p('{anexo'+str(i)+'Instruccion}',12,center=True,space=24,line=1.5)
 if i<=10:
  note='El modelo original está en la siguiente página. Usa las copias necesarias para los alumnos del equipo y conserva todas las hojas completadas.'
  if i==9:note='Imprime el formulario de la siguiente página y las copias que necesites por alumno. Llénalo a mano; después gestiona la firma y el sello y escanea las hojas.'
  if i==6:note='El modelo de la siguiente página es solo para identificar la certificación. La Unidad debe emitir el documento oficial después de la aprobación.'
 else:note='Adjunta aquí el documento que ya presentaste, revisado y con sus respaldos. No se incluye una plantilla nueva para este anexo.' if i<13 else 'Adjunta tus solicitudes y respuestas, si existen. Si no solicitaste prórroga, retira esta hoja sin agregar un documento.'
 p(note,12,center=True,space=22,line=1.5)
 p('Retira esta hoja de indicaciones antes de entregar. Incorpora el documento completado con las firmas y sellos que correspondan.',12,True,True,space=0,line=1.5)
 if i<=10:append_original(*form_ranges[i-1],expected[i-1])
head('Asesorías y evidencias del proyecto')
p('HOJA DE INDICACIONES',11,True,space=20)
p('{respaldosIndicaciones}',12,space=18,line=1.5)
p('Respaldos adicionales declarados',12,True);paras('anexos')
p('Retira esta hoja e incorpora tus evidencias antes de las prórrogas, si las hay. Revisa el índice y entrega un único PDF con los anexos completos y legibles.',12,True,line=1.5)
# Copied drawings must have unique Word drawing identifiers.
for n,node in enumerate(d.element.body.iter(qn('wp:docPr')),1):node.set('id',str(n))
# No automatic external updates. Page numbers are internal Word fields.
for el in list(d.settings.element.findall(qn('w:updateFields'))):d.settings.element.remove(el)
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'false');d.settings.element.append(u)
q=d.sections[0].footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');q._p.append(f)
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='Memoria de servicio social';d.core_properties.comments=''
out=R/'documentos/plantillas/memoria.docx';d.save(out)
(R/'js/memoria-plantilla.js').write_text('// Institutional memory template embedded for static/offline generation.\nwindow.SIGPSMemoriaPlantilla="'+base64.b64encode(out.read_bytes()).decode()+'";\n')
